import datetime
import os
import logging

from dateutil import parser
from openpyxl import Workbook, load_workbook

from django.conf import settings

from seahub.project.models import Projects
from seahub.knowledge_base.models import KnowledgeBaseViews
from seahub.project.seadb_api import SeaDBAPI
from seahub.project.view_utils import view_data_2_sql
from seahub.project.constants import KNOWLEDGE_BASE_EXPORT_DISPLAY_ALL_COLUMNS


logger = logging.getLogger('seaqa_io')


def filter_display_columns(columns):
    return [c for c in (columns or []) if c.get('name') in KNOWLEDGE_BASE_EXPORT_DISPLAY_ALL_COLUMNS]


def convert_kb_view_to_excel(project_uuid, view_id, username):
    target_dir = os.path.join(settings.TEMP_EXPORT_VIEW_DIR, str(project_uuid))
    os.makedirs(target_dir, exist_ok=True)

    # Fetch project name and view name
    project = Projects.objects.get_project_by_uuid(project_uuid)
    project_name = project.name if project else 'project'

    view = KnowledgeBaseViews.objects.get_view(project_uuid, view_id)
    view_name = (view or {}).get('name') or 'view'

    # Fetch columns metadata from SeaDB
    seadb_api = SeaDBAPI(username)
    metadata = seadb_api.get_base_metadata(project_uuid)

    tables_metadata = metadata.get('tables') or []
    kb_table_name = 'knowledge_base'
    table_meta = next((t for t in tables_metadata if t.get('name') == kb_table_name), {})
    columns = table_meta.get('columns', [])
    display_columns = filter_display_columns(columns) or columns

    # Build SQL and query rows from SeaDB
    sql = view_data_2_sql(kb_table_name, display_columns, view or {}, username, 0, 1000)
    res = seadb_api.query_rows(project_uuid, sql, convert_keys=False)
    rows = res.get('results', [])

    # Write XLSX
    column_name_map = {
        'title': 'Title',
        'content': 'Content',
        'creator': 'Creator',
        'created_time': 'Created time',
        'last_modifier': 'Last modifier',
        'modified_time': 'Last modified time',
    }

    excel_name = f'{project_name}_knowledge_base_{view_name}.xlsx'
    target_path = os.path.join(target_dir, excel_name)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(view_name or 'view')

    header = [column_name_map.get(c.get('name'), c.get('name')) for c in (display_columns or [])]
    header_names = [c.get('name') for c in (display_columns or [])]
    header_keys = [c.get('key') for c in (display_columns or [])]

    if header:
        ws.append(header)

    for row in (rows or []):
        row_data = []
        for col_name, col_key in zip(header_names, header_keys):
            val = row.get(col_key)
            if col_name in ('created_time', 'modified_time') and val:
                try:
                    dt = parser.parse(str(val))
                    val = dt.strftime('%Y-%m-%d %H:%M:%S')
                except Exception:
                    pass
            row_data.append(val)
        ws.append(row_data)

    wb.save(target_path)
    return {}


def preview_import_kb_from_excel(project_uuid, file_name, limit=20):
    temp_dir = os.path.join(settings.TEMP_EXPORT_VIEW_DIR, str(project_uuid))
    file_path = os.path.join(temp_dir, file_name)

    wb = load_workbook(file_path, read_only=True)
    ws = wb.active

    headers = []
    rows_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        raise Exception('Empty file')

    column_name_map = {
        'title': 'Title',
        'content': 'Content',
    }
    header_to_key = {v.lower(): k for k, v in column_name_map.items()}

    column_indexes = {}
    for idx, header in enumerate(headers):
        if not header:
            continue
        h = str(header).strip().lower()
        key = header_to_key.get(h)
        if key:
            column_indexes[key] = idx

    if 'title' not in column_indexes or 'content' not in column_indexes:
        return {'file_name': file_name, 'preview_rows': [], 'total_rows': 0}

    preview_rows = []
    total_rows = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        row_obj = {}
        for key, idx in column_indexes.items():
            if idx >= len(row):
                continue
            value = row[idx]
            if value is None:
                continue
            row_obj[key] = value
        if 'title' in row_obj and 'content' in row_obj:
            total_rows += 1
            if len(preview_rows) < limit:
                preview_rows.append({'title': row_obj.get('title'), 'content': row_obj.get('content')})

    return {'file_name': file_name, 'preview_rows': preview_rows, 'total_rows': total_rows}


def import_kb_from_excel(project_uuid, username, file_name):
    temp_dir = os.path.join(settings.TEMP_EXPORT_VIEW_DIR, str(project_uuid))
    file_path = os.path.join(temp_dir, file_name)

    wb = load_workbook(file_path, read_only=True)
    ws = wb.active

    headers = []
    rows_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        raise Exception('Empty file')

    column_name_map = {
        'title': 'Title',
        'content': 'Content',
    }
    header_to_key = {v.lower(): k for k, v in column_name_map.items()}

    column_indexes = {}
    for idx, header in enumerate(headers):
        if not header:
            continue
        h = str(header).strip().lower()
        key = header_to_key.get(h)
        if key:
            column_indexes[key] = idx

    if 'title' not in column_indexes or 'content' not in column_indexes:
        return {}

    rows_to_insert = []
    seadb_api = SeaDBAPI(username)
    kb_table_name = 'knowledge_base'
    now = datetime.datetime.now(datetime.timezone.utc).isoformat()

    try:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            row_obj = {}
            for key, idx in column_indexes.items():
                if idx >= len(row):
                    continue
                value = row[idx]
                if value is None:
                    continue
                row_obj[key] = value
            if 'title' in row_obj and 'content' in row_obj:
                row_obj['creator'] = username
                row_obj['last_modifier'] = username
                row_obj['created_time'] = now
                row_obj['modified_time'] = now
                row_obj['deleted'] = False
                rows_to_insert.append(row_obj)

            if len(rows_to_insert) >= 1000:
                seadb_api.insert_rows(project_uuid, kb_table_name, rows_to_insert)
                rows_to_insert = []

        if rows_to_insert:
            seadb_api.insert_rows(project_uuid, kb_table_name, rows_to_insert)
    finally:
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
        except Exception as e:
            logger.warning(e)

    return {}
